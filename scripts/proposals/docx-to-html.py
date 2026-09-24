"""Read proposal DOCX content and create printable HTML for Chromium PDF export."""

from __future__ import annotations

import base64
import html
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "temp" / "proposal-tools"))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

OUT = ROOT / "temp" / "proposals" / "skylimit"


def image_uri(paragraph, document):
    blips = paragraph._p.xpath(".//a:blip")
    if not blips:
        return None
    rel = blips[0].get(qn("r:embed"))
    part = document.part.related_parts[rel]
    return "data:image/png;base64," + base64.b64encode(part.blob).decode("ascii")


def para_html(paragraph):
    content = html.escape(paragraph.text)
    if not content.strip():
        return ""
    sizes = [round(run.font.size.pt, 1) for run in paragraph.runs if run.font.size]
    size = max(sizes, default=10.5)
    if size >= 25:
        tag, cls = "h1", "title"
    elif size >= 13:
        tag, cls = "h2", "subhead"
    elif size >= 11:
        tag, cls = "p", "lead"
    elif size <= 8:
        tag, cls = "p", "caption"
    elif size <= 9:
        tag, cls = "p", "eyebrow" if any(r.bold for r in paragraph.runs) else "smalltext"
    else:
        tag, cls = "p", "bodytext"
    if "http" in content:
        import re
        content = re.sub(r"(https?://[^\s<]+)", r'<a href="\1">\1</a>', content)
    return f'<{tag} class="{cls}">{content}</{tag}>'


def table_html(table):
    rows = table.rows
    if len(rows) == 1 and len(rows[0].cells) == 3:
        cells = []
        for cell in rows[0].cells:
            parts = [html.escape(p.text) for p in cell.paragraphs if p.text.strip()]
            cells.append(f'<div><b>{parts[0]}</b><span>{parts[1]}</span></div>')
        return '<div class="stats">' + "".join(cells) + '</div>'
    if len(rows) == 1 and len(rows[0].cells) == 1:
        parts = [html.escape(p.text) for p in rows[0].cells[0].paragraphs if p.text.strip()]
        return f'<div class="card"><b>{parts[0]}</b><span>{parts[1]}</span></div>'
    result = []
    for row in rows:
        number = html.escape(row.cells[0].text.strip())
        parts = [html.escape(p.text) for p in row.cells[1].paragraphs if p.text.strip()]
        result.append(f'<div class="step"><strong>{number}</strong><div><b>{parts[0]}</b><span>{parts[1]}</span></div></div>')
    return '<div class="steps">' + "".join(result) + '</div>'


def convert(stem):
    document = Document(OUT / (stem + ".docx"))
    cover = None
    pages = [[]]
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            image = image_uri(paragraph, document)
            if image and cover is None:
                cover = image
                continue
            if paragraph._p.xpath('.//w:br[@w:type="page"]'):
                pages.append([])
                continue
            if image:
                pages[-1].append(f'<img class="content-shot" src="{image}">')
            else:
                rendered = para_html(paragraph)
                if rendered:
                    pages[-1].append(rendered)
        elif element.tag == qn("w:tbl"):
            pages[-1].append(table_html(Table(element, document)))
    assert cover and len(pages) == 5, (bool(cover), len(pages))
    style = """
    @page{size:8.5in 11in;margin:0}*{box-sizing:border-box}html,body{margin:0;padding:0;font-family:Arial,Helvetica,sans-serif;color:#071a33}
    .page{width:8.5in;height:11in;page-break-after:always;position:relative;overflow:hidden;background:white}
    .cover img{width:8.5in;height:11in;display:block}.body-page{padding:.70in .77in .64in}
    .content{height:9.63in;overflow:hidden}.eyebrow{font-size:8.5pt;letter-spacing:1.3px;color:#169bb4;font-weight:700;text-transform:uppercase;margin:0 0 11pt}
    h1{font-size:27pt;line-height:1.08;letter-spacing:-.5pt;margin:0 0 10pt}h2{font-size:13.5pt;margin:16pt 0 7pt}.lead{font-size:11.3pt;line-height:1.35;color:#4c6070;margin:0 0 19pt}
    .bodytext{font-size:10.25pt;line-height:1.34;margin:0 0 10pt}.caption,.smalltext{font-size:8pt;color:#4c6070;line-height:1.25;margin:5pt 0 11pt}.caption{letter-spacing:.25pt}
    .stats{display:flex;gap:4px;margin:5pt 0 6pt}.stats>div{flex:1;background:#071a33;color:#eef9fa;padding:15pt 12pt;min-height:76pt}.stats b{display:block;font-size:23pt;color:#7de6f1}.stats span{display:block;font-size:9pt;line-height:1.2;margin-top:5pt}
    .card{background:#e9f6f8;padding:11pt 13pt;margin:0 0 8pt;break-inside:avoid}.card b,.card span{display:block}.card b{font-size:11pt;margin-bottom:4pt}.card span{font-size:9.8pt;color:#4c6070;line-height:1.25}
    .steps{margin:5pt 0 14pt}.step{display:flex;border-bottom:1px solid #c8dfe3;padding:7pt 3pt;gap:14pt}.step strong{font-size:10pt;color:#169bb4;min-width:20pt}.step b,.step span{display:block}.step b{font-size:10pt}.step span{font-size:9.3pt;color:#4c6070;line-height:1.24;margin-top:3pt}
    .content-shot{width:100%;height:auto;display:block;margin:7pt 0 2pt;border:1px solid #d5e4e7;border-radius:5px}
    .footer{position:absolute;right:.77in;bottom:.29in;color:#708492;font-size:7.2pt;letter-spacing:.4px}
    a{color:#168ca5;text-decoration:underline;word-break:break-all}
    """
    page_markup = '<section class="page cover"><img src="' + cover + '"></section>'
    for body_blocks in pages:
        page_markup += '<section class="page body-page"><div class="content">' + "\n".join(body_blocks) + '</div><div class="footer">360VISION  •  PRIVATE PROPOSAL  •  AHMED ABDELAZIZ</div></section>'
    output = OUT / (stem + ".html")
    output.write_text(f'<!doctype html><html><head><meta charset="utf-8"><style>{style}</style></head><body>{page_markup}</body></html>', encoding="utf8")
    print(output)


if __name__ == "__main__":
    convert("360Vision-SkyLimit-Proposal")
    convert("360Vision-Real-Estate-Proposal-Template")
