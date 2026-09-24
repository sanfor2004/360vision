"""Build editable, client-specific and reusable 360Vision proposal documents.

Install once into the ignored tools directory:
  python -m pip install python-docx --target temp/proposal-tools
Then run:
  python scripts/proposals/build-docx.py
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "temp" / "proposal-tools"))

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION_START  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

OUT = ROOT / "temp" / "proposals" / "skylimit"
MARKETING = ROOT / "public" / "markting"
NAVY = RGBColor(7, 26, 51)
CYAN = RGBColor(22, 155, 180)
SLATE = RGBColor(76, 96, 112)
PALE = "E9F6F8"
WHITE = "FFFFFF"


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "C8DFE3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("bottom",):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)
        borders.append(node)
    tc_pr.append(borders)


def set_cell_margin(cell, top=150, start=180, bottom=150, end=180):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        elt = OxmlElement(f"w:{side}")
        elt.set(qn("w:w"), str(value))
        elt.set(qn("w:type"), "dxa")
        mar.append(elt)
    tc_pr.append(mar)


def line(doc, text="", style=None, size=None, color=None, bold=False, after=8, before=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def heading(doc, label: str, title: str, lead: str):
    line(doc, label.upper(), size=8.5, color=CYAN, bold=True, after=10)
    line(doc, title, size=27, color=NAVY, bold=True, after=10)
    line(doc, lead, size=11.5, color=SLATE, after=20)


def subhead(doc, text: str, after=6):
    line(doc, text, size=14, color=NAVY, bold=True, after=after, before=8)


def body(doc, text: str, after=10):
    return line(doc, text, size=10.5, color=NAVY, after=after)


def card(doc, title: str, text: str, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.95)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, fill)
    set_cell_margin(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = NAVY
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = SLATE
    line(doc, after=7)


def stats(doc):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    for col in table.columns:
        col.width = Inches(2.30)
    data = [
        ("33%", "put floor plans first"),
        ("20%", "put virtual tours first"),
        ("75%", "of REALTORS® use social media"),
    ]
    for i, (value, label) in enumerate(data):
        cell = table.cell(0, i)
        shade(cell, "071A33")
        set_cell_margin(cell, 200, 180, 200, 180)
        p = cell.paragraphs[0]
        r = p.add_run(value)
        r.bold = True
        r.font.size = Pt(24)
        r.font.color.rgb = RGBColor(125, 230, 241)
        p = cell.add_paragraph()
        r = p.add_run(label)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(233, 246, 248)
    line(doc, "Sources: Zillow 2025 prospective-buyer research; NAR 2025 REALTORS® Technology Survey.", size=8, color=SLATE, after=18, before=6)


def steps(doc, items):
    table = doc.add_table(rows=len(items), cols=2)
    table.autofit = False
    table.columns[0].width = Inches(.64)
    table.columns[1].width = Inches(6.25)
    for i, (title, detail) in enumerate(items):
        a, b = table.rows[i].cells
        for cell in (a, b):
            set_cell_border(cell)
            set_cell_margin(cell, 100, 90, 100, 90)
        p = a.paragraphs[0]
        r = p.add_run(f"{i + 1:02d}")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = CYAN
        p = b.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = NAVY
        p = b.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(detail)
        r.font.size = Pt(9.5)
        r.font.color.rgb = SLATE
    line(doc, after=12)


def picture(doc, filename, width=6.95, caption=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(filename), width=Inches(width))
    if caption:
        line(doc, caption, size=8, color=SLATE, after=12)


def page(doc):
    doc.add_page_break()


def add_hyperlink(paragraph, label: str, url: str):
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "169BB4")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def source(doc, name, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{name} — ")
    r.font.size = Pt(8.5)
    r.font.color.rgb = SLATE
    add_hyperlink(p, url, url)


def init_doc(cover: Path, generic: bool):
    doc = Document()
    doc.core_properties.title = "360Vision | Real-Estate Growth Proposal"
    doc.core_properties.subject = "Funded real-estate pilot proposal"
    doc.core_properties.author = "Ahmed Abdelaziz"
    first = doc.sections[0]
    first.page_width = Inches(8.5)
    first.page_height = Inches(11)
    first.top_margin = Inches(.03)
    first.bottom_margin = Inches(.03)
    first.left_margin = Inches(.03)
    first.right_margin = Inches(.03)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().add_picture(str(cover), width=Inches(8.42))
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.top_margin = Inches(.72)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.78)
    section.right_margin = Inches(.78)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.line_spacing = 1.12
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("360VISION  •  PRIVATE PROPOSAL  •  AHMED ABDELAZIZ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = SLATE
    return doc


def build(generic=False):
    cover = OUT / ("cover-generic-letter.png" if generic else "cover-letter.png")
    doc = init_doc(cover, generic)
    audience = "a prospective real-estate partner" if generic else "SkyLimit"

    # Page 2 — decision context.
    heading(doc, "01 / Opportunity", "Make space part of the sales conversation", "People can inspect a property online before deciding whether it deserves an in-person visit. Agents need a clear way to turn that attention into a useful next conversation.")
    stats(doc)
    body(doc, "A good tour lets a buyer understand room relationships, move between floors, and return to a space that matters. For an agent, it becomes a more useful listing asset when it sits beside professional images and a straightforward way to request a showing.")
    subhead(doc, "The opportunity is a complete property presentation")
    card(doc, "For buyers", "Explore the layout at their own pace through connected 360° rooms, clear labels, and an optional floor-plan map.")
    card(doc, "For agents", "Present a property consistently, reuse selected moments in short social clips, and invite an interested buyer to take the next step.")
    body(doc, "This is an opportunity to test, not a promise that tours alone sell homes faster. Zillow provides a free basic tour option, so the proposed offer must earn its place through presentation quality, service, and a useful sales handoff.", after=0)

    # Page 3 — the buyer-facing workflow.
    page(doc)
    heading(doc, "02 / Customer workflow", "From property capture to qualified conversation", "An agent should be able to commission or provide imagery, approve one polished tour, and use its best moments across listing and social channels.")
    steps(doc, [
        ("Capture the space", "Use professional equirectangular 360° photography through a local partner, or approved panoramas supplied by the customer."),
        ("Build the tour", "Name rooms, set the opening view, connect scenes with hotspots, and add optional floor plans and floor points."),
        ("Publish the experience", "In the funded build, create a customer-facing hosted tour suitable for sharing from a property page or approved listing channels."),
        ("Start the conversation", "In the funded build, offer an explicit inquiry or showing request and explore a consent-based handoff into the customer’s sales workflow."),
    ])
    picture(doc, MARKETING / "screenshots" / "viewer-desktop.png", 6.9, "REAL PRODUCT CAPTURE  ·  The viewer keeps room navigation and the optional property map separate.")

    # Page 4 — what exists and what needs funding.
    page(doc)
    heading(doc, "03 / Product proof", "A working foundation, ready for a commercial pilot", "360Vision already runs as a local, single-owner authoring studio and viewer. Its current experience can be demonstrated now; customer delivery features must be built and validated.")
    picture(doc, MARKETING / "screenshots" / "studio-desktop.png", 6.9, "REAL PRODUCT CAPTURE  ·  Studio authors panorama scenes, hotspots, opening camera views, and optional floor maps.")
    subhead(doc, "Working today")
    body(doc, "Scene and hotspot authoring; connected room navigation; optional multi-floor plans; mobile viewer layouts; local JSON autosave; panorama processing; and short demo footage created from the real app.")
    subhead(doc, "Funded build scope")
    body(doc, "A hosted read-only tour, reliable media delivery, a branded property presentation, a clear inquiry action, deployment operations, and pilot-specific validation. Any SkyLimitApp connection would be designed with SkyLimit after confirming its current API and data requirements.")
    card(doc, "Clear boundary", "The current local app does not provide public hosting, client accounts, lead capture, or a live SkyLimitApp integration. None of those are claimed as completed in this proposal.")

    # Page 5 — partner fit.
    page(doc)
    heading(doc, "04 / Partner fit", f"Why this fits {audience}", "The strongest version connects a useful visual experience to a sales process without making an agent manage another complicated tool.")
    if generic:
        body(doc, "A brokerage or property team can use 360Vision to make an individual listing easier to understand, then connect an interested viewer to the team’s existing follow-up process. The first pilot should prove that workflow with one real property.")
    else:
        body(doc, "SkyLimitApp publicly presents lead storage, follow-up reminders, calling, SMS, email, and reporting in one marketing workflow. A property tour could become a new source of buyer interest for that workflow. This is a proposed product direction, not an existing integration. [1]")
    steps(doc, [
        ("Show", "An agent sends a hosted 360° tour with a clear floor-plan view and focused property story."),
        ("Signal", "A buyer actively requests a showing or asks a question. Passive viewing is not treated as consent to contact."),
        ("Follow up", "The agent responds through their existing process. A future integration could route an opted-in inquiry into SkyLimitApp."),
    ])
    subhead(doc, "A service-led entry, then a software business")
    body(doc, "Start with a paid property pilot and optional photography coordination. If agents value the result and repeat the workflow, test a repeatable tour-production package and software pricing. Let real customer demand determine whether a larger platform or dedicated brand is justified.")
    subhead(doc, "Where it can grow")
    body(doc, "First: residential brokerages and listing teams. Next: photographers as production partners, rental and property managers, commercial spaces, and stores that benefit from remote walkthroughs. These are adjacent markets to test, not proven customers.")
    card(doc, "Three commercial levers to validate", "A paid tour for each listing; optional professional capture and social-media production; and, if repeat usage justifies it, a team subscription for authoring and hosting. The pilot will test willingness to pay before any pricing model is fixed.")

    # Page 6 — paid ask and sources.
    page(doc)
    heading(doc, "05 / Decision", "Fund one focused pilot. Learn before scaling.", "I’m proposing a meeting to define a paid build with one launch property and one real estate customer or partner willing to evaluate it.")
    card(doc, "Pilot outcome", "Deliver one shareable property tour, an optional map when a plan is available, a small set of social cutdowns, and a documented inquiry path. Professional 360° capture is separately scoped through a local partner or supplied by the client.")
    card(doc, "What we will learn", "Can an agent publish and explain the tour easily? Can buyers find and understand rooms? Will a buyer choose to request a conversation? Will the team pay for a repeatable version?")
    card(doc, "Commercial next step", "Meet this week to choose the pilot property, buyer, hosting approach, schedule, responsibilities, and a funded scope with agreed milestones. Pricing and ownership terms follow that scoping conversation.")
    subhead(doc, "Orel — let’s discuss the build this week" if not generic else "Let’s discuss a pilot this week")
    body(doc, "I built 360Vision and can lead the technical work. The attached short film shows the real product in motion. I’d value 20 minutes to show you the prototype and decide whether this is a useful real-estate direction to fund.")
    line(doc, "Prepared for: Property team" if generic else "Prepared for: Orel David, SkyLimit LLC", size=8.5, color=SLATE, after=6)
    p = line(doc, "Ahmed Abdelaziz  |  Creator of 360Vision  |  ", size=9, color=SLATE, after=9)
    add_hyperlink(p, "sanfor2004.github.io", "https://sanfor2004.github.io")
    line(doc, "Research and company sources", size=8.5, color=CYAN, bold=True, after=4, before=5)
    if not generic:
        source(doc, "[1] SkyLimitApp product and pricing", "https://skylimitapp.com/")
    source(doc, "Zillow, Prospective Buyers: Consumer Housing Trends 2025", "https://www.zillow.com/research/prospective-buyers-consumer-housing-trends-2025/")
    source(doc, "NAR, 2025 REALTORS® Technology Survey", "https://www.nar.realtor/press-releases/realtors-embrace-ai-digital-tools-to-enhance-client-service-nar-survey-finds")
    line(doc, "Visual provenance: screenshots and motion are captures of the working Cedar House demo; the cover's lifestyle background is illustrative.", size=8, color=SLATE, after=0)

    filename = "360Vision-SkyLimit-Proposal.docx" if not generic else "360Vision-Real-Estate-Proposal-Template.docx"
    doc.save(OUT / filename)
    print(OUT / filename)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build(False)
    build(True)
    (OUT / "whatsapp-message.txt").write_text(
        "FIRST MESSAGE\n\n"
        "Hi Orel, I hope you’re well. After my work at SkyLimit, I built 360Vision—a working interactive property-tour studio. "
        "I see a real-estate product opportunity that could connect better property presentation with the kind of sales follow-up SkyLimit already knows well. "
        "I made a short video and a concise proposal for you. Could we meet for 20 minutes this week to look at a funded pilot? "
        "I’d lead the build and use one property to test the offer with agents.\n\nAhmed\n\n"
        "FOLLOW-UP (IF NEEDED AFTER A FEW BUSINESS DAYS)\n\n"
        "Hi Orel, following up on the 360Vision preview I sent. The main idea is a focused paid pilot: one property tour, one real customer test, and a clear decision on whether the real-estate vertical is worth building further. "
        "If it interests you, I can show the working prototype in 20 minutes this week.\n",
        encoding="utf8",
    )
    (OUT / "README.txt").write_text(
        "360VISION / SKYLIMIT PROPOSAL HANDOFF\n\n"
        "Send by WhatsApp in this order: first message, whatsapp-cover.png, 360vision-skylimit-preview.mp4, and 360Vision-SkyLimit-Proposal.pdf. "
        "The DOCX is the editable version for revisions. Do not send both Word and PDF unless requested.\n\n"
        "The personalized files are private local deliverables. The generic Word and PDF template can be adapted for another brokerage or property team. "
        "No messages have been sent and no files have been uploaded.\n\n"
        "Visual provenance: the screen images and original footage come from the working Cedar House demo. "
        "The lifestyle background on the cover is illustrative. Hosting, inquiry capture, and SkyLimit integration are proposed build work, not current product features.\n",
        encoding="utf8",
    )
